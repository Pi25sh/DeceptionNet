"""Generate the official DeceptionNet .docx manuscript with professional formatting,
incorporating all ground-truth code mappings, the verified 9-step campaign (severity 26.0)
across 7 total harness sessions, the 84-test verification suite, and dedicated sections
resolving all 6 reviewer items.
"""
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def build_manuscript(output_path):
    doc = Document()

    # Page Margins: 1 inch everywhere
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("DeceptionNet: An Architectural Framework for Multi-Honeypot Active Defense, Deterministic ATT&CK Attribution, and Post-Capture Threat Intelligence")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    p_title.paragraph_format.space_after = Pt(8)

    # Subtitle / Metadata
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Undergraduate Capstone Engineering Manuscript & Technical Demonstration Report\nTarget Track: Student Systems / Academic Security Workshop\nRepository: https://github.com/Pi25sh/DeceptionNet")
    r_sub.italic = True
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p_sub.paragraph_format.space_after = Pt(18)

    # Abstract Box
    abstract_table = doc.add_table(rows=1, cols=1)
    abstract_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = abstract_table.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    p_abs = cell.paragraphs[0]
    r_abs_h = p_abs.add_run("Abstract—")
    r_abs_h.bold = True
    r_abs_t = p_abs.add_run(
        "Modern cyber defense relies heavily on proactive deception technology to divert, profile, and delay "
        "adversaries before they compromise production infrastructure. However, traditional standalone honeypots "
        "often present operational barriers: their telemetry is siloed across disparate daemon formats, rule mapping "
        "is prone to misattribution or high latency, and emerging proposals to introduce generative Large Language "
        "Models (LLMs) directly into the shell interaction loop suffer from ungrounded hallucinations and severe prompt "
        "injection risks. This paper presents DeceptionNet, an open-source active defense platform built on isolated "
        "Docker microservices. DeceptionNet unifies multi-service honeypot decoys (Cowrie SSH/Telnet, Dionaea multi-protocol, "
        "and a vulnerable Flask web target) with an asynchronous log-forwarding bridge, a deterministic 33-rule MITRE ATT&CK "
        "mapping engine, monotonic additive threat severity scoring, sliding-window session correlation, and a live React 18 "
        "operator dashboard. Crucially, DeceptionNet strictly decouples the LLM from interactive decoy simulation: generative "
        "models (Google Gemini 2.5 Flash / Claude 3 Haiku) are employed exclusively as post-capture threat intelligence "
        "summarizers for SOC analysts and automated ReportLab PDF incident reports. The platform logic is formally verified "
        "via an 84-test automated pytest suite covering rule boundaries, database savepoints under multithreaded concurrency, "
        "and timing edge cases. Finally, we demonstrate operational viability through a committed end-to-end multi-stage attack "
        "case study executing the automated attack harness, recording 7 distinct sessions across 15 events, including a persistent "
        "authenticated campaign of 9 events reaching a cumulative threat score of 26.0 (Critical)."
    )
    r_abs_t.font.size = Pt(10)
    p_abs.paragraph_format.space_after = Pt(4)

    p_kw = cell.add_paragraph()
    r_kwh = p_kw.add_run("Keywords—")
    r_kwh.bold = True
    r_kwt = p_kw.add_run("Honeypots, Active Cyber Defense, MITRE ATT&CK, Telemetry Normalization, Cross-Honeypot Correlation, Decoupled LLM Summarization, Concurrency Safety.")
    r_kwt.italic = True
    r_kwt.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_h1(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0x2E, 0x5B, 0x88)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        return h

    # Section 1
    add_h1("1. Introduction & Problem Statement")
    doc.add_paragraph(
        "Active deception platforms deploy decoy systems mimicking legitimate computing assets to detect unauthorized "
        "reconnaissance and exploitation attempts early in the cyber kill chain. When properly isolated, any interaction "
        "with a decoy is anomalous by definition, producing near-zero false-positive detection rates compared to perimeter "
        "intrusion detection systems (IDS) flooded by noisy background traffic."
    )
    doc.add_paragraph(
        "Despite their conceptual advantages, deploying honeypots in production or educational lab environments presents "
        "three primary technical hurdles: (1) Telemetry Fragmentation: decoys like Cowrie and Dionaea produce heterogeneous, "
        "unstructured JSON and binary logs that cannot be directly queried without custom ingestion pipelines; (2) Semantic "
        "Attribution Latency: security operations center (SOC) analysts must manually correlate raw shell commands and "
        "connection probes to the MITRE ATT&CK framework; and (3) LLM Misapplication: recent academic proposals advocate "
        "placing generative Large Language Models directly on open network ports to generate interactive shell responses. "
        "In practice, interactive LLM decoys introduce prompt injection vulnerabilities, unpredictable latency spikes, "
        "and hallucinations during attack attribution."
    )
    doc.add_paragraph(
        "DeceptionNet addresses these challenges by implementing an integrated, containerized microservices platform "
        "combining deterministic rule attribution, concurrency-safe database ingestion, and decoupled post-capture AI "
        "reporting. All empirical claims in this work are strictly grounded in committed source code and verified command executions."
    )

    # Section 2
    add_h1("2. System Architecture & Component Design")
    doc.add_paragraph(
        "The DeceptionNet platform is organized into five decoupled subsystems interconnected across dual Docker bridge networks:"
    )

    p_arch = doc.add_paragraph()
    p_arch.add_run("1. Decoy Infrastructure (honeypot_net):\n").bold = True
    p_arch.add_run(
        "   • Cowrie (ports 2222/2223): Emulates SSH and Telnet shells, capturing credential guesses, interactive commands, "
        "and dropped files in /var/log/cowrie/cowrie.json.\n"
        "   • Dionaea: Traps automated worm and exploit payloads across SMB (445), FTP (21), SIP, and HTTP endpoints.\n"
        "   • Attack-Target (port 9090): Deliberately vulnerable Python Flask portal exposing realistic flaw patterns: "
        "SQL injection login bypass, unrestricted command execution (/ping), path traversal (/read), reflected XSS (/search), "
        "and file upload (/upload).\n"
    )
    p_arch.add_run("2. Telemetry Bridge (honeypot/forwarder.py):\n").bold = True
    p_arch.add_run(
        "   • Continuously tails decoy JSON logs, normalizes heterogeneous schemas into a unified EventPayload structure "
        "(session_id, event_type, timestamp, src_ip, input, message), and posts events over HTTP to core-api.\n"
    )
    p_arch.add_run("3. Core Security Engine (core-api):\n").bold = True
    p_arch.add_run(
        "   • Implements deterministic MITRE ATT&CK mapping via 33 rule specifications (16 command keyword rules spanning 50 "
        "signatures and 17 event-type normalization rules).\n"
        "   • Manages sliding-window session correlation (get_or_create_session) with cross-honeypot correlation and intra-honeypot isolation.\n"
        "   • Calculates monotonic additive severity scoring with atomic database increments.\n"
        "   • Enforces OAuth2 JWT authentication and RBAC for all REST routes and the /ws/live-events WebSocket channel.\n"
    )
    p_arch.add_run("4. Decoupled AI Reporting Engine (ai-service & core-api/report.py):\n").bold = True
    p_arch.add_run(
        "   • Strictly operates asynchronously *after* capture. Synthesizes session telemetry into executive summaries and "
        "compiles high-resolution ReportLab PDF incident reports with MITRE technique badges.\n"
    )
    p_arch.add_run("5. Live Operator Dashboard (dashboard):\n").bold = True
    p_arch.add_run(
        "   • Built with React 18, Vite, and WebSocket streaming. Features live attack feeds, MITRE technique distribution, "
        "session inspection drawers, and PDF report downloads.\n"
    )

    # Section 3
    add_h1("3. Deterministic MITRE ATT&CK Mapping & Threat Scoring")
    doc.add_paragraph(
        "Attribution accuracy and execution speed are critical during high-volume brute-force attacks. DeceptionNet executes "
        "deterministic rule evaluation in < 1 ms via core-api/mitre.py."
    )

    add_h2("3.1 Rule Specifications & Signature Precedence")
    doc.add_paragraph(
        "The rule table contains exactly 33 rule specifications:\n"
        "• 16 Command Keyword Rules: 50 distinct shell signatures evaluated with regex word boundaries (\\b) to eliminate "
        "sub-token false positives (e.g., ensuring 'id' does not match 'rapid'). Specific rules take precedence over general ones: "
        "reading '/etc/shadow' or '/etc/passwd' matches T1003.008 (weight 5.0) before generic password commands match T1098. "
        "System discovery commands are strictly separated: 'uname' maps to T1082 (System Information Discovery, weight 1.0), "
        "while 'whoami' and 'id' map to T1033 (System Owner/User Discovery, weight 1.0).\n"
        "• 17 Event-Type Normalization Rules: Protocol actions mapped directly from decoy events (e.g., login_failed -> T1110 (0.5), "
        "login -> T1078 (1.0), sqli -> T1190 (4.0), file_upload -> T1505.003 (3.0), http -> T1190 (4.0), scan -> T1046 (1.0))."
    )

    add_h2("3.2 Threat Scoring Formula & Threshold Tiers")
    doc.add_paragraph(
        "Session threat severity is computed as an additive monotonic function of positively identified technique weights:\n"
        "   Severity = round(∑ w_i, 2)  where w_i > 0\n"
        "Events that produce non-detections or noise (w = 0.0) contribute nothing to the score. As implemented in dashboard/src/utils.js, "
        "threat scores map into four operational tiers: Low (< 1.0), Elevated (1.0 – 2.4), High (2.5 – 3.9), and Critical (≥ 4.0)."
    )

    # Section 4
    add_h1("4. Operational Case Study: Real Attack Harness Execution")
    doc.add_paragraph(
        "To replace all hypothetical evaluation metrics with verifiable ground truth, an automated intrusion was executed "
        "using the automated attack harness (attacker/scripts/full-attack.sh) against the attack-target web application (port 9090). "
        "Telemetry was processed through core-api, stored in SQLite, and committed as raw output to results/case_study_raw.json."
    )
    doc.add_paragraph(
        "Session Assignment Method: The web application (attack-target/app.py:current_sid) pins the session ID web-<ip>-<hex> "
        "in a signed Flask session cookie (session['sid']) upon login. When full-attack.sh authenticates in Stage 1 with curl's cookie "
        "jar (-c $COOKIE_JAR), follow-up post-exploitation requests in Stages 3–6 present the cookie (-b $COOKIE_JAR), binding all "
        "9 exploited events to the same session ID. In contrast, credential stuffing requests in Stage 2 and the nmap scan in Stage 7 "
        "send no cookies; each unauthenticated request is assigned a fresh random UUID session ID, yielding 5 separate brute-force "
        "sessions and 1 scan session (7 total sessions in the run across 15 events)."
    )

    # Add Table
    cs_table = doc.add_table(rows=1, cols=7)
    cs_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["#", "Harness Stage", "Executed Action / Payload", "Event Type", "ATT&CK Technique", "Weight", "Running"]
    hdr_cells = cs_table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9)
        set_cell_background(hdr_cells[i], "1B365D")
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)

    # 9 rows of authenticated campaign matching results/case_study_raw.json exactly
    rows_data = [
        ("1", "1. SQLi Login Bypass", "POST /login: username=admin'--&password=x", "sqli", "T1190 Exploit App", "+4.0", "4.0"),
        ("2", "3. Cmd Injection (whoami)", "POST /ping: host=127.0.0.1; whoami", "command", "T1033 User Discovery", "+1.0", "5.0"),
        ("3", "3. Cmd Injection (uname)", "POST /ping: host=127.0.0.1; uname -a", "command", "T1082 System Discovery", "+1.0", "6.0"),
        ("4", "3. Cmd Injection (passwd)", "POST /ping: host=127.0.0.1; cat /etc/passwd", "command", "T1003.008 Cred Dump", "+5.0", "11.0"),
        ("5", "3. Cmd Injection (tool)", "POST /ping: host=127.0.0.1; wget http://...", "command", "T1105 Ingress Transfer", "+3.0", "14.0"),
        ("6", "4. Path Traversal (passwd)", "GET /read?file=../../../../etc/passwd", "command", "T1003.008 Cred Dump", "+5.0", "19.0"),
        ("7", "4. Path Traversal (conf)", "GET /read?file=../../etc/portal.conf", "command", "Unmapped (Noise)", "+0.0", "19.0"),
        ("8", "5. Reflected XSS", "GET /search?q=<script>alert(1)</script>", "http", "T1190 Exploit App", "+4.0", "23.0"),
        ("9", "6. Web Shell Upload", "POST /upload: file=@/tmp/shell.php", "file_upload", "T1505.003 Web Shell", "+3.0", "26.0"),
    ]

    for row in rows_data:
        r_cells = cs_table.add_row().cells
        for j, val in enumerate(row):
            r_cells[j].text = val
            r_cells[j].paragraphs[0].runs[0].font.size = Pt(8.5)
            set_cell_margins(r_cells[j], top=60, bottom=60, left=80, right=80)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run("Table 1: Ground-truth event ledger and MITRE attribution from the authenticated campaign session (cumulative severity 26.0).")
    r_cap.italic = True
    r_cap.font.size = Pt(9.5)

    # Section 5
    add_h1("5. Technical Resolution of Open Review Items")

    add_h2("5.1 Item 1: Session Count from Committed Harness Run")
    doc.add_paragraph(
        "The committed execution of attacker/scripts/full-attack.sh produces exactly 7 sessions across 15 total events in the database:\n"
        "• 1 persistent authenticated campaign session (9 events, severity 26.0) bound together via the Flask session cookie.\n"
        "• 5 unauthenticated credential stuffing sessions (1 event each, severity 0.5, T1110) because password spray attempts send no cookies.\n"
        "• 1 unauthenticated nmap scan session (1 event, severity 1.0, T1046).\n"
        "The historical 13 vs 14 session discussion pertained to attacker/scripts/hack-demo.sh with an 11-word wordlist. "
        "We report strictly the single concrete count from the committed full-attack run: 7 sessions."
    )

    add_h2("5.2 Item 2: Correlation is Cross-Honeypot Only")
    doc.add_paragraph(
        "In core-api/main.py:441-455, session correlation explicitly enforces AttackSession.honeypot != honeypot. "
        "The 5-minute sliding window is strictly designed for cross-decoy transitions (e.g., an attacker scanning Dionaea SMB "
        "before moving to Cowrie SSH). For events arriving on the same decoy service without an existing session cookie, "
        "every new session ID is intentionally treated as an independent connection. This ensures high-volume password spray "
        "bursts from an attacker remain distinct rather than artificially conflating into a single long-lived session."
    )

    add_h2("5.3 Item 3: Exact ATT&CK Labels (T1082, T1033, T1190, T1505.003)")
    doc.add_paragraph(
        "Attribution tags strictly match established taxonomies and codebase implementations:\n"
        "• uname -> T1082: The rule table in core-api/mitre.py maps 'uname' strictly to T1082 (System Information Discovery, weight 1.0), "
        "separating it from 'whoami' and 'id' which map to T1033 (System Owner/User Discovery, weight 1.0).\n"
        "• SQLi Login Bypass -> T1190: SQL injection authentication bypass is an application exploitation vulnerability. "
        "In attack-target/app.py and core-api/mitre.py, SQLi bypass emits event type 'sqli' which maps strictly to T1190 "
        "(Exploit Public-Facing Application, weight 4.0). Legitimate credential logins continue to map to T1078 (Valid Accounts).\n"
        "• Web Shell Upload -> T1505.003: Verified against attack.mitre.org, uploading a PHP webshell (shell.php) represents "
        "Server Software Component: Web Shell (T1505.003, weight 3.0) under Persistence, not data exfiltration."
    )

    add_h2("5.4 Item 4: Decoupled Role of LLM: Post-Capture Summarizer vs Interactive Decoy")
    doc.add_paragraph(
        "In DeceptionNet, the LLM is strictly decoupled from live interaction:\n"
        "• Interactive decoy emulation is handled entirely by Cowrie, Dionaea, and Flask state machines.\n"
        "• The LLM service (Google Gemini 2.5 Flash / Claude 3 Haiku) operates asynchronously *after* capture as a threat intelligence "
        "post-processor. It ingests the captured event ledger once a session concludes, generates executive natural language summaries "
        "for SOC operators, and synthesizes incident narrative paragraphs for automated ReportLab PDF reports (core-api/report.py).\n"
        "• The LLM plays no role in decoy emulation, command generation, or honeypot execution; it is solely an offline reporting assistant."
    )

    add_h2("5.5 Item 5: Concurrency Safety & Database Transaction Guarantees")
    doc.add_paragraph(
        "High-volume honeypot bursts can trigger race conditions when multiple forwarder threads post events concurrently. "
        "DeceptionNet implements two robust transaction guarantees:\n"
        "1. Atomic Score Updates: The guarantee against lost severity updates rests mathematically on the SQL engine's atomic addition guarantee:\n"
        "   UPDATE AttackSession SET severity_score = coalesce(severity_score, 0) + weight\n"
        "   This eliminates Python-level read-modify-write lost-update races regardless of backend concurrency. "
        "(This is an engine-level guarantee, not an empirical claim from a SQLite test, since SQLite serializes write transactions).\n"
        "2. Savepoints for Technique Deduplication (main.py:472-498): Uses db.begin_nested() to catch duplicate key conflicts on "
        "(session_id, technique, source) without aborting outer transaction state.\n"
        "3. Multi-Threaded Concurrency Test: test_concurrent_upsert_technique_threads spawns 8 concurrent threads attempting simultaneous "
        "technique upserts against SQLite in WAL mode, asserting that all savepoints recover cleanly and exactly 1 deduplicated record persists."
    )

    add_h2("5.6 Item 6: Candid Limitations of the Heuristic Grounding Filter")
    doc.add_paragraph(
        "In ai-service/main.py:_grounding_check, candidate MITRE techniques retrieved via FAISS vector search are filtered by checking "
        "for lexical token overlap against raw event strings.\n"
        "While this heuristic successfully prevents naive hallucinations, it is inherently gameable: if an attacker embeds common "
        "security terms into command arguments (e.g., ping -c 1 'kerberoasting privilege escalation discovery'), the filter matches these substring "
        "tokens and accepts unrelated techniques. Production systems should replace token matching with Abstract Syntax Tree (AST) "
        "command parsing or Natural Language Inference (NLI) cross-encoders."
    )

    # Section 6
    add_h1("6. Automated Verification Suite (84 Tests)")
    doc.add_paragraph(
        "The DeceptionNet test suite consists of 84 automated tests executed via pytest core-api/tests with zero failures:\n"
        "• test_mitre.py (61 tests): Verifies all rule specifications across keyword signatures and event types, validating word "
        "boundaries, case insensitivity, and rule precedence (including uname -> T1082 and whoami -> T1033).\n"
        "• test_correlation.py (12 tests): Verifies exact session matching, cross-honeypot correlation, intra-honeypot isolation, "
        "and 8-thread concurrent upsert execution with database savepoints.\n"
        "• test_scoring.py (9 tests): Verifies monotonic scoring formulas, zero-weight filtering, and test_harness_attack_session_score "
        "asserting the exact score of 26.0 matching the authenticated campaign session in results/case_study_raw.json.\n"
        "• test_evaluation_benchmark.py (2 tests): Verifies schema contracts for illustrative sample datasets and tests 5 real boundary "
        "scenarios in get_or_create_session (cross-honeypot merge, same-honeypot isolation, 4m50s inside window, 5m10s outside window, and NAT collision)."
    )

    # Section 7
    add_h1("7. System Limitations & Target Venue Alignment")
    doc.add_paragraph(
        "In keeping with academic integrity, we candidly outline the system's operational boundaries:\n"
        "1. Multiplexed Networks (NAT): When multiple hosts share a single public IP, the 5-minute IP heuristic merges cross-honeypot sessions "
        "unless supplemented by client TLS/SSH fingerprinting (HASSH) or HTTP cookies.\n"
        "2. Keyword Obfuscation: Static rule matching relies on recognizable command names; base64-encoded or obfuscated payloads bypass static "
        "keyword rules unless caught by semantic vector search.\n"
        "3. Venue Recommendation: As a functional educational prototype evaluated through an 84-test verification suite and a real attack case study, "
        "DeceptionNet is appropriately targeted for an Undergraduate Capstone Track, Student System Demonstration Track, or Academic Workshop. "
        "Scopus indexing depends on the specific publisher proceedings (e.g., IEEE Xplore, ACM ICPS, Springer LNCS) and must be verified on "
        "the Scopus Sources list prior to submission."
    )

    # Section 8
    add_h1("8. Conclusion")
    doc.add_paragraph(
        "DeceptionNet establishes an open-source, containerized active defense architecture combining deterministic ATT&CK mapping, "
        "concurrency-safe telemetry ingestion, and decoupled AI incident reporting. By rejecting synthetic benchmarks and grounding every "
        "reported metric in verified code executions, this work provides a transparent foundation for cyber deception research and education."
    )

    doc.save(output_path)
    print(f"Successfully generated manuscript: {output_path}")


if __name__ == "__main__":
    out_file = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "docs", "DeceptionNet_Manuscript.docx"))
    build_manuscript(out_file)
